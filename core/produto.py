# -*- coding: utf-8 -*-
"""Figuras e pacote .zip do diagnóstico — município, UC ou TI (docs/10 §1).

``Diagnostico_<MUN|UC|TI>_<Territorio>_<ano1>_<ano2>.zip``:
    relatorio_diagnostico_<…>.pdf, estatisticas_<…>.xlsx, grade_10x10km_<…>.geojson
    (+ grade_<tam>x<tam>km_<…>.geojson proporcional),
    dados_<MUN|UC|TI>_<…>.csv, ranking_municipios/ucs/tis.csv, metadados.json,
    imagens/*.png (gráficos, mapas, estatísticas no modelo INPE e, no município, o cenário atual),
    estatisticas/*.csv (tabelas no modelo INPE), nowcast/*.csv (município),
    risco_de_recorrencia/*.kml + .geojson (4 classes de recorrência, por quartis, dissolvidas)

O pacote é montado em memória e baixado pelo navegador.
"""

import io
import json
import zipfile
from dataclasses import dataclass

import matplotlib.pyplot as plt
import pandas as pd
from shapely.geometry import mapping

from core.config import APP_NOME, APP_VERSAO, SIGLA_TIPO, TAMANHO_CELULA_KM, TIPO_MUNICIPIO
from core.diagnostico import MESES, DiagnosticoMunicipio, ranking_simples
from core.estatisticas import (
    EstatisticasMunicipio, extremos_por_coluna, notas, tabela_acumulado_exibicao, tabela_comparativo_exibicao,
    tabela_diaria_exibicao, tabela_mensal_exibicao,
)
from core.grade import grade_qgis, kml_risco, nota_grade, risco_de_recorrencia
from core.graficos import grafico_evolucao_tendencia, grafico_top_estado
from core.graficos_estatisticas import figuras_estatisticas, grafico_nowcast, grafico_nowcast_semanal
from core.mapas import mapa_densidade, mapa_dispersao, mapa_recentes, mapa_recorrencia, mapa_trimestral
from core.relatorio_pdf import gerar_pdf
from core.territorios import area_km2
from core.texto import neutralizar_formulas, nome_arquivo

DPI_PNG = 180
DIAS_MAPA_RELATORIO = 7          # mapa do cenário atual no PDF/pacote
# figura → nome do PNG no pacote (sem o sufixo _<alvo>.png)
NOMES_PNG = {
    "top_estado": "grafico_top5_estado",
    "evolucao": "grafico_evolucao_historica",
    "sazonalidade": "grafico_sazonalidade",
    "mapa_dispersao": "mapa_1_densidade_ano_final",
    "mapa_densidade": "mapa_2_densidade_acumulada",
    "mapa_recorrencia": "mapa_3_recorrencia",
    "mapa_trimestral": "mapa_4_sazonalidade_trimestral",
    "comparativo_mensal": "estatisticas_comparativo_mensal",
    "comparativo_semanal": "estatisticas_comparativo_semanal",
    "comparativo_semestre1": "estatisticas_comparativo_1o_semestre",
    "comparativo_semestre2": "estatisticas_comparativo_2o_semestre",
    "acumulado_mensal": "estatisticas_acumulado_mensal",
    "nowcast_diario": "cenario_atual_ultimos_30_dias",
    "nowcast_semanal": "cenario_atual_ultimas_52_semanas",
    "nowcast_mapa": "cenario_atual_mapa_focos_recentes",
}
COR_MINIMO = "FFD54F"     # destaque do INPE: menor valor da coluna em amarelo, maior em vermelho
COR_MAXIMO = "E53935"


@dataclass
class Pacote:
    nome: str
    conteudo: bytes
    pdf: bytes | None = None


# ------------------------------------------------------------------ Figuras

def figuras_territorio(diag, est: EstatisticasMunicipio, nowcast=None, dias_mapa: int = DIAS_MAPA_RELATORIO) -> dict:
    """Todas as figuras do diagnóstico (as do cenário atual só com ``nowcast``)."""
    municipio = isinstance(diag, DiagnosticoMunicipio)
    focos_ano = diag.focos[diag.focos["ano"] == diag.ano_fim]
    areas = diag.areas if municipio else None
    figuras = {
        "evolucao": grafico_evolucao_tendencia(diag.serie_anual),
        "mapa_dispersao": mapa_dispersao(focos_ano, diag.poligono, diag.municipio, diag.ano_fim, areas),
        "mapa_densidade": mapa_densidade(diag.focos, diag.poligono, diag.municipio, diag.ano_inicio, diag.ano_fim),
        "mapa_recorrencia": mapa_recorrencia(diag.focos, diag.poligono, diag.ano_inicio, diag.ano_fim),
        "mapa_trimestral": mapa_trimestral(diag.focos, diag.poligono, diag.municipio, diag.ano_inicio, diag.ano_fim),
        **figuras_estatisticas(est),
    }
    if municipio:
        figuras["top_estado"] = grafico_top_estado(diag.top_acumulado, diag.municipio)
    if nowcast is not None:
        figuras["nowcast_diario"] = grafico_nowcast(nowcast)
        figuras["nowcast_semanal"] = grafico_nowcast_semanal(nowcast)
        figuras["nowcast_mapa"] = mapa_recentes(nowcast.recentes(dias_mapa), diag.poligono, diag.municipio,
                                                nowcast.ultimo_dia, dias_mapa, areas)
    return figuras


def fechar_figuras(figuras: dict):
    for figura in figuras.values():
        if figura is not None:
            plt.close(figura)


def png(figura, dpi: int = DPI_PNG) -> bytes:
    buffer = io.BytesIO()
    figura.savefig(buffer, format="png", bbox_inches="tight", dpi=dpi)
    return buffer.getvalue()


# ----------------------------------------------------------------- Tabelas

def tabela_sazonalidade(est: EstatisticasMunicipio) -> pd.DataFrame:
    """Média mensal dos anos do período (sem o de referência), p10, p90 e o ano de referência."""
    p = est.parametros
    meses = est.mensal[list(range(1, 13))]
    base = meses.loc[[a for a in est.anos_estatistica if a in meses.index]]
    return pd.DataFrame({
        "Mês": MESES,
        "Média": base.mean().round(1).to_numpy(),
        f"p{p.percentil_inferior}": base.quantile(p.percentil_inferior / 100).round(1).to_numpy(),
        f"p{p.percentil_superior}": base.quantile(p.percentil_superior / 100).round(1).to_numpy(),
        str(est.ano_referencia): meses.loc[est.ano_referencia].to_numpy() if est.ano_referencia in meses.index
        else None,
    })


def tabelas_estatisticas(est: EstatisticasMunicipio) -> dict[str, pd.DataFrame]:
    """Aba/arquivo → tabela das estatísticas no modelo INPE (com o índice como 1ª coluna)."""
    return {
        "Focos_Mes_a_Mes": tabela_mensal_exibicao(est),
        "Focos_Dia_a_Dia": tabela_diaria_exibicao(est),
        "Comparativo_Mensal": tabela_comparativo_exibicao(est, "mensal").set_index("Mês"),
        "Comparativo_Semanal": tabela_comparativo_exibicao(est, "semanal").set_index("Semana"),
        "Comparativo_1o_Semestre": tabela_comparativo_exibicao(est, "semestre1").set_index("Semana do semestre"),
        "Comparativo_2o_Semestre": tabela_comparativo_exibicao(est, "semestre2").set_index("Semana do semestre"),
        "Acumulado_Mensal": tabela_acumulado_exibicao(est),
    }


def tabelas_nowcast(nowcast) -> dict[str, pd.DataFrame]:
    serie = nowcast.serie.assign(data=[d.strftime("%d/%m/%Y") for d in nowcast.serie["data"]]).rename(columns={
        "data": "Data", "focos": "Focos", "media": "Média do dia (anos anteriores)",
        "maximo": "Máximo do dia (anos anteriores)", "media_janela": "Média ±3 dias (anos anteriores)"})
    recentes = nowcast.focos.assign(data=[d.strftime("%d/%m/%Y") for d in nowcast.focos["data"]]).rename(
        columns={"data": "Data", "dias_atras": "Dias antes do último", "idade": "Idade", "lat": "Latitude",
                 "lon": "Longitude"})
    semanal = nowcast.semanal.assign(**{c: [d.strftime("%d/%m/%Y") for d in nowcast.semanal[c]]
                                         for c in ("inicio", "fim")}).rename(columns={
        "inicio": "Início", "fim": "Fim", "focos": "Focos", "media": "Média da semana (anos anteriores)",
        "maximo": "Máximo da semana (anos anteriores)"})
    indicadores = pd.DataFrame(list(nowcast.indicadores.items()), columns=["Indicador", "Valor"])
    return {"Nowcast_Ultimos_30_Dias": serie, "Nowcast_Ultimas_52_Semanas": semanal,
            "Nowcast_Focos_Recentes": recentes, "Nowcast_Indicadores": indicadores}


def _destacar_extremos(planilha, tabela: pd.DataFrame, linhas: list[str]):
    """Pinta o menor (amarelo) e o maior (vermelho) valor de cada coluna, como no site do INPE."""
    from openpyxl.styles import Font, PatternFill

    posicao_linha = {nome: i + 2 for i, nome in enumerate(tabela.index)}
    for j, (minimos, maximos) in enumerate(extremos_por_coluna(tabela, linhas).values(), start=2):
        for nome in minimos:
            celula = planilha.cell(row=posicao_linha[nome], column=j)
            celula.fill, celula.font = PatternFill("solid", fgColor=COR_MINIMO), Font(bold=True)
        for nome in maximos:
            celula = planilha.cell(row=posicao_linha[nome], column=j)
            celula.fill, celula.font = PatternFill("solid", fgColor=COR_MAXIMO), Font(bold=True, color="FFFFFF")


def planilha_estatisticas(diag, est: EstatisticasMunicipio, nowcast=None) -> bytes:
    """Excel: evolução anual, sazonalidade, estatísticas no modelo INPE e, se houver, o cenário atual."""
    evolucao = pd.DataFrame({"Ano": list(diag.serie_anual), "Focos": list(diag.serie_anual.values())})
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        evolucao.to_excel(escritor, sheet_name="Evolucao_Historica", index=False)
        tabela_sazonalidade(est).to_excel(escritor, sheet_name="Sazonalidade_Mensal", index=False)
        for aba, tabela in tabelas_estatisticas(est).items():
            tabela.to_excel(escritor, sheet_name=aba)
        anos = [str(a) for a in est.anos_estatistica]   # o ano de referência não é destacado
        _destacar_extremos(escritor.sheets["Focos_Mes_a_Mes"], tabela_mensal_exibicao(est), anos)
        pd.DataFrame({"Notas": notas(est)}).to_excel(escritor, sheet_name="Notas_Estatisticas", index=False)
        if nowcast is not None:
            for aba, tabela in tabelas_nowcast(nowcast).items():
                tabela.to_excel(escritor, sheet_name=aba, index=False)
        neutralizar_formulas(escritor.book)
    return buffer.getvalue()


def metadados(nome: str, tipo: str, ano_inicio: int, ano_fim: int, total_focos: int,
              extras: dict | None = None) -> bytes:
    return json.dumps({
        "nome_alvo": nome,
        "tipo_analise": tipo,
        "ano_inicio": ano_inicio,
        "ano_fim": ano_fim,
        "total_focos": total_focos,
        **(extras or {}),
        "gerado_por": f"{APP_NOME} v{APP_VERSAO}",
    }, ensure_ascii=False, indent=2).encode("utf-8")


# ------------------------------------------------------------------- Pacote

def _montar_zip(arquivos: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, conteudo in arquivos.items():
            z.writestr(nome, conteudo)
    return buffer.getvalue()


PASTA_RISCO = "risco_de_recorrencia"


def arquivos_risco(diag, alvo: str) -> dict[str, bytes]:
    """Pasta de risco de recorrência: 4 KML (baixo, moderado, alto, muito alto), um polígono dissolvido cada,
    e um GeoJSON com as 4 classes juntas."""
    risco = risco_de_recorrencia(diag.focos, diag.poligono)
    nota = nota_grade(risco["grade"]) if risco["grade"] is not None else ""
    periodo = f"{diag.ano_inicio}–{diag.ano_fim}"
    areas = [area_km2(c["geometria"]) if c["geometria"] is not None else None for c in risco["classes"]]
    arquivos = {
        f"{PASTA_RISCO}/risco_{n}_{classe['codigo']}_{alvo}.kml": kml_risco(
            classe, diag.municipio, periodo, risco["quartis"], nota, area).encode("utf-8")
        for n, (classe, area) in enumerate(zip(risco["classes"], areas), start=1)
    }
    # As 4 classes juntas, para o QGIS (P26); classes vazias ficam de fora
    q = risco["quartis"] or (None, None, None)
    feicoes = [{
        "type": "Feature",
        "properties": {"ordem": n, "classe": classe["codigo"], "rotulo": classe["rotulo"], "cor_hex": classe["cor"],
                       "anos_min": classe["anos"][0], "anos_max": classe["anos"][1], "celulas": classe["celulas"],
                       "area_km2": round(area, 1) if area else None, "q1": q[0], "q2": q[1], "q3": q[2]},
        "geometry": mapping(classe["geometria"]),
    } for n, (classe, area) in enumerate(zip(risco["classes"], areas), start=1) if classe["geometria"] is not None]
    arquivos[f"{PASTA_RISCO}/risco_recorrencia_classes_{alvo}.geojson"] = json.dumps({
        "type": "FeatureCollection", "name": f"risco_recorrencia_{alvo}",
        "descricao": f"{diag.municipio} — {periodo}. Classes pelos quartis da recorrência das células com foco. {nota}",
        "features": feicoes}, ensure_ascii=False).encode("utf-8")
    return arquivos


def pacote(diag, figuras: dict, est: EstatisticasMunicipio, nowcast=None) -> Pacote:
    """Pacote completo do diagnóstico (município, UC ou TI); ``figuras`` = ``figuras_territorio``."""
    municipio = isinstance(diag, DiagnosticoMunicipio)
    tipo = TIPO_MUNICIPIO if municipio else diag.tipo
    sigla = SIGLA_TIPO[tipo]
    alvo = nome_arquivo(diag.municipio)
    periodo = f"{diag.ano_inicio}_{diag.ano_fim}"
    pdf = gerar_pdf(diag, figuras, est, nowcast)
    extras = ({"uf": diag.sigla_uf, "tipo_regiao": diag.tipo_regiao, "regiao": diag.regiao} if municipio
              else {"uf_predominante": diag.sigla_uf})
    if nowcast is not None:
        extras["nowcast_ate"] = nowcast.ultimo_dia.isoformat()
    # P31: grade de 10 km sempre (comparável entre territórios) + a proporcional (a do mapa 3 e do risco)
    grades = {TAMANHO_CELULA_KM: grade_qgis(diag.focos, diag.poligono, TAMANHO_CELULA_KM)}
    proporcional = grade_qgis(diag.focos, diag.poligono)
    grades.setdefault(proporcional.get("tamanho_celula_km", TAMANHO_CELULA_KM), proporcional)
    arquivos = {
        f"relatorio_diagnostico_{alvo}_{periodo}.pdf": pdf,
        f"estatisticas_{alvo}_{periodo}.xlsx": planilha_estatisticas(diag, est, nowcast),
        f"dados_{sigla}_{alvo}.csv": diag.focos.to_csv(index=False).encode("utf-8"),
        "metadados.json": metadados(diag.municipio, tipo, diag.ano_inicio, diag.ano_fim, len(diag.focos), extras),
    }
    for km, grade in grades.items():
        arquivos[f"grade_{km:g}x{km:g}km_{alvo}.geojson"] = json.dumps(grade, indent=2, ensure_ascii=False).encode()
    focos_rankings = diag.focos_estado if municipio else diag.focos
    for coluna, nome in (("NM_MUN", "ranking_municipios.csv"), ("UC", "ranking_ucs.csv"), ("TI", "ranking_tis.csv")):
        arquivos[nome] = ranking_simples(focos_rankings, coluna, diag.fim_dados).to_csv(index=False).encode("utf-8")
    for chave, figura in figuras.items():
        if figura is not None:
            arquivos[f"imagens/{NOMES_PNG[chave]}_{alvo}.png"] = png(figura)
    for i, (nome, tabela) in enumerate(tabelas_estatisticas(est).items(), start=1):
        arquivos[f"estatisticas/{i:02d}_{nome}.csv"] = tabela.to_csv().encode("utf-8")
    if nowcast is not None:
        for nome, tabela in tabelas_nowcast(nowcast).items():
            arquivos[f"nowcast/{nome}.csv"] = tabela.to_csv(index=False).encode("utf-8")
    arquivos.update(arquivos_risco(diag, alvo))
    return Pacote(f"Diagnostico_{sigla}_{alvo}_{periodo}.zip", _montar_zip(arquivos), pdf)
