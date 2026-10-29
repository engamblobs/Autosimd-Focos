# -*- coding: utf-8 -*-
"""Grade 10 km, kernel, figuras, PDF, Excel e pacote .zip (docs/10), para município e área protegida."""

import io
import json
import math
import re
import zipfile
from datetime import date

import numpy as np
import pandas as pd
import pytest
from shapely.geometry import box

from core.config import TIPO_UC
from core.diagnostico import diagnosticar_area, diagnosticar_municipio
from core.estatisticas import cobertura_base, estatisticas_territorio
from core.focos import carregar_focos, focos_do_municipio
from core.grade import contagem, grade_do_territorio, grade_qgis, nota_grade, recorrencia, risco_de_recorrencia
from core.kernel import densidade_kernel
from core.mapas import MAX_CRUZES, mapa_densidade, mapa_recentes
from core.nowcast import nowcast_territorio
from core.produto import NOMES_PNG, fechar_figuras, figuras_territorio, pacote, planilha_estatisticas
from core.territorios import (
    areas_protegidas_no_territorio, areas_protegidas_padrao, estados_disponiveis, geometria_area_protegida,
    geometria_municipio,
)

UC_TRIUNFO = "ÁREA DE PROTEÇÃO AMBIENTAL TRIUNFO DO XINGU (1039)"


@pytest.fixture(scope="module")
def df():
    return carregar_focos()


@pytest.fixture(scope="module")
def municipio(df):
    poligono = geometria_municipio(estados_disponiveis()["Pará"], "NOVO PROGRESSO")
    areas = areas_protegidas_no_territorio(poligono, areas_protegidas_padrao())
    diag = diagnosticar_municipio(df, "Pará", "NOVO PROGRESSO", 2023, 2024, poligono, areas)
    cobertura = cobertura_base(df)
    est = estatisticas_territorio(diag.focos, diag.municipio, 2023, 2024, cobertura)
    nowcast = nowcast_territorio(focos_do_municipio(df, "Pará", "NOVO PROGRESSO"), diag.municipio,
                                 date(2024, 12, 31), cobertura[0])
    figuras = figuras_territorio(diag, est, nowcast)
    yield diag, figuras, pacote(diag, figuras, est, nowcast), est
    fechar_figuras(figuras)


@pytest.fixture(scope="module")
def area(df):
    diag = diagnosticar_area(df, TIPO_UC, UC_TRIUNFO, 2023, 2024, geometria_area_protegida("UC", UC_TRIUNFO))
    est = estatisticas_territorio(diag.focos, diag.nome, 2023, 2024, cobertura_base(df))
    figuras = figuras_territorio(diag, est)
    yield diag, figuras, pacote(diag, figuras, est)
    fechar_figuras(figuras)


def conteudo(pacote_) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(pacote_.conteudo))


# ------------------------------------------------------------- Grade e kernel

def test_grade_ajustada_a_latitude():
    """P8: 10 km nos dois eixos, em qualquer latitude."""
    norte = grade_do_territorio((-55.0, -1.0, -54.0, 1.0), 10)        # Equador
    sul = grade_do_territorio((-53.0, -31.0, -52.0, -29.0), 10)       # −30°
    assert norte.passo_y == pytest.approx(sul.passo_y)
    assert norte.passo_x == pytest.approx(10 / 111.32, rel=1e-3)
    assert sul.passo_x == pytest.approx(10 / (111.32 * math.cos(math.radians(-30))), rel=1e-3)
    km_x = sul.passo_x * 111.32 * math.cos(math.radians(-30))
    km_y = sul.passo_y * 110.574
    assert km_x == pytest.approx(10, rel=1e-3) and km_y == pytest.approx(10, rel=1e-3)
    assert "ajustada à latitude" in nota_grade(sul)


def test_contagem_e_recorrencia():
    grade = grade_do_territorio((0.0, 0.0, 0.3, 0.3), 10)
    lons, lats, anos = [0.01, 0.02, 0.01, 0.25], [0.01, 0.01, 0.02, 0.25], [2020, 2020, 2021, 2021]
    total = contagem(grade, lons, lats)
    assert total.sum() == 4 and total.max() == 3
    assert recorrencia(grade, lons, lats, anos).max() == 2       # célula com fogo em 2020 e 2021
    assert contagem(grade, [], []).sum() == 0


def test_grade_proporcional_ao_territorio():
    """Célula = 10 km ÷ 1, 2, 4 ou 8, para o lado maior ter ~51 células (São Félix do Xingu-PA com 10 km)."""
    def celula(largura_km, altura_km, lat=-5.0):
        dx = largura_km / (111.32 * math.cos(math.radians(lat)))
        return grade_do_territorio((-50.0, lat - altura_km / 221.148, -50.0 + dx, lat + altura_km / 221.148)).tamanho_km
    assert celula(304, 511, -7.5) == 10            # São Félix do Xingu: referência
    assert celula(440, 736) == 10                  # maior que a referência: não passa de 10 km
    assert celula(296, 153) == 5                   # Marabá: largo, conta o lado maior
    assert celula(26, 50) == 1.25                  # Maracanã
    assert celula(10, 16) == 1.25                  # Marituba: mínimo de 1,25 km
    grade = grade_do_territorio((-48.0, -1.0, -47.76, -0.55))
    assert grade.tamanho_km == 1.25 and grade.passo_y == pytest.approx(1.25 / 110.574)
    assert 30 <= max(grade.forma) <= 60
    assert "1,25 km × 1,25 km (10 km ÷ 8" in nota_grade(grade)


def test_kernel_em_territorio_menor_que_o_filtro():
    """APA Pedregulho (~16 × 8 km): o raster tem menos células que a gaussiana (antes: erro de formato)."""
    territorio = box(-47.26, -23.262, -47.116, -23.186)
    k = densidade_kernel([-47.2, -47.15], [-23.22, -23.2], territorio.bounds, territorio)
    assert k.densidade.shape == (len(k.y) - 1, len(k.x) - 1)
    assert np.nansum(k.densidade) > 0


def test_kernel_conserva_os_focos():
    territorio = box(-52.0, -6.0, -51.0, -5.0)
    rng = np.random.default_rng(1)
    lons, lats = rng.uniform(-51.8, -51.2, 200), rng.uniform(-5.8, -5.2, 200)
    k = densidade_kernel(lons, lats, territorio.bounds)
    assert np.nansum(k.densidade) * k.celula_km ** 2 == pytest.approx(200, rel=0.01)
    recortado = densidade_kernel(lons, lats, territorio.bounds, territorio)
    assert np.isnan(recortado.densidade).any()                   # fora do limite = vazio
    assert np.nansum(recortado.densidade) * k.celula_km ** 2 == pytest.approx(200, rel=0.05)


def test_grade_qgis(municipio):
    diag = municipio[0]
    geo = grade_qgis(diag.focos, diag.poligono)
    focos = [f["properties"]["focos_acumulados"] for f in geo["features"]]
    assert sum(focos) == len(diag.focos)                   # todo foco cai numa célula
    assert all(1 <= f["properties"]["recorrencia_anos"] <= 2 for f in geo["features"])
    assert "ajustada à latitude" in geo["descricao"]
    json.dumps(geo)
    assert grade_qgis(diag.focos.iloc[0:0])["features"] == []


def test_mapa_recentes_sem_fundo(municipio):
    """Sem o fundo de satélite (testes / sem internet) o mapa sai com aviso."""
    figura = municipio[1]["nowcast_mapa"]
    assert any("indisponível" in t.get_text() for t in figura.texts)
    vazio = mapa_recentes(pd.DataFrame(columns=["lon", "lat", "idade"]), None, "X", date(2024, 1, 1), 7)
    fechar_figuras({"x": vazio})


# ------------------------------------------------------------------ Pacotes

def test_figuras_municipio(municipio):
    figuras = municipio[1]
    # a amostra (2023–2024) tem um só ano de referência: as anomalias precisam de pelo menos 3 (docs/07 §12)
    assert set(figuras) == {k for k in NOMES_PNG if not k.startswith("anomalia_")}
    assert all(f is not None for f in figuras.values())


def test_pacote_municipio(municipio):
    diag, _, pacote_, _ = municipio
    assert pacote_.nome == "Diagnostico_MUN_NOVO_PROGRESSO_2023_2024.zip"
    nomes = set(conteudo(pacote_).namelist())
    assert {
        "relatorio_diagnostico_NOVO_PROGRESSO_2023_2024.pdf", "estatisticas_NOVO_PROGRESSO_2023_2024.xlsx",
        "dados_MUN_NOVO_PROGRESSO.csv", "metadados.json", "grade_10x10km_NOVO_PROGRESSO.geojson", "grade_5x5km_NOVO_PROGRESSO.geojson",
        "ranking_municipios.csv", "ranking_ucs.csv", "ranking_tis.csv",
        "nowcast/Nowcast_Ultimos_30_Dias.csv", "nowcast/Nowcast_Ultimas_13_Semanas.csv", "nowcast/Nowcast_Focos_Recentes.csv",
    } <= nomes
    assert {f"imagens/{n}_NOVO_PROGRESSO.png" for k, n in NOMES_PNG.items() if not k.startswith("anomalia_")} <= nomes
    assert {f"estatisticas/0{i}_{n}.csv" for i, n in enumerate(
        ["Focos_Mes_a_Mes", "Focos_Dia_a_Dia", "Comparativo_Mensal", "Comparativo_Semanal",
         "Comparativo_1o_Semestre", "Comparativo_2o_Semestre", "Acumulado_Mensal"], start=1)} <= nomes
    z = conteudo(pacote_)
    metadados = json.loads(z.read("metadados.json"))
    assert metadados["total_focos"] == 5923 and metadados["tipo_analise"] == "Estado / Município"
    assert (metadados["regiao"], metadados["tipo_regiao"]) == ("Região Tapajós", "Região de Integração")
    assert metadados["nowcast_ate"] == "2024-12-31"
    ranking = pd.read_csv(io.BytesIO(z.read("ranking_municipios.csv")))
    assert len(ranking) == 144 and ranking["Focos de Calor"].sum() == len(diag.focos_estado)


def test_pdf_do_municipio(municipio):
    pdf = municipio[2].pdf
    assert pdf.startswith(b"%PDF")
    assert int(re.search(rb"/Count (\d+)", pdf).group(1)) >= 9
    assert 300_000 < len(pdf) < 8_000_000


def test_planilha(municipio):
    diag, _, pacote_, est = municipio
    planilhas = pd.read_excel(io.BytesIO(conteudo(pacote_).read("estatisticas_NOVO_PROGRESSO_2023_2024.xlsx")),
                              sheet_name=None)
    assert list(planilhas) == ["Evolucao_Historica", "Sazonalidade_Mensal", "Focos_Mes_a_Mes", "Focos_Dia_a_Dia",
                               "Comparativo_Mensal", "Comparativo_Semanal", "Comparativo_1o_Semestre",
                               "Comparativo_2o_Semestre", "Acumulado_Mensal", "Notas_Estatisticas",
                               "Nowcast_Ultimos_30_Dias", "Nowcast_Ultimas_13_Semanas", "Nowcast_Focos_Recentes", "Nowcast_Indicadores"]
    assert planilhas["Evolucao_Historica"]["Focos"].tolist() == [1136, 4787]
    sazonal = planilhas["Sazonalidade_Mensal"]
    assert list(sazonal.columns) == ["Mês", "Média", "p10", "p90", "2024"]
    assert sazonal["Média"].tolist() == est.mensal.loc[2023, list(range(1, 13))].tolist()   # só 2023 na média
    from openpyxl import load_workbook
    aba = load_workbook(io.BytesIO(planilha_estatisticas(diag, est)))["Focos_Mes_a_Mes"]
    cores = {aba.cell(row=2, column=j).fill.fgColor.rgb for j in range(2, 15)}
    assert any(c and c.endswith("E53935") for c in cores)          # máximo da coluna em vermelho (linha 2023)


def test_pacote_da_uc_tem_pdf_e_planilha(area):
    """P7: UC/TI também com PDF e Excel."""
    diag, figuras, pacote_ = area
    assert "top_estado" not in figuras and "nowcast_mapa" not in figuras
    assert pacote_.pdf.startswith(b"%PDF")
    assert pacote_.nome == "Diagnostico_UC_AREA_DE_PROTECAO_AMBIENTAL_TRIUNFO_DO_XINGU_1039_2023_2024.zip"
    nomes = set(conteudo(pacote_).namelist())
    alvo = "AREA_DE_PROTECAO_AMBIENTAL_TRIUNFO_DO_XINGU_1039"
    assert {f"relatorio_diagnostico_{alvo}_2023_2024.pdf", f"estatisticas_{alvo}_2023_2024.xlsx",
            f"dados_UC_{alvo}.csv"} <= nomes
    assert not any(n.startswith("nowcast/") for n in nomes)
    ranking = pd.read_csv(io.BytesIO(conteudo(pacote_).read("ranking_ucs.csv")))
    assert ranking["UC"].tolist() == [UC_TRIUNFO] and ranking["Focos de Calor"].iloc[0] == diag.focos_periodo


# ------------------------------------------------- Risco de recorrência e cruzes

def test_risco_de_recorrencia_por_quartis():
    rng = np.random.default_rng(3)
    territorio = box(-52.0, -6.0, -51.0, -5.0)
    linhas = []
    for celula in range(40):                                     # 40 pontos fixos com 1 a 10 anos de fogo
        lon, lat = rng.uniform(-51.9, -51.1), rng.uniform(-5.9, -5.1)
        for ano in range(2000, 2000 + 1 + celula % 10):
            linhas.append({"lon": lon, "lat": lat, "ano": ano})
    focos = pd.DataFrame(linhas)
    risco = risco_de_recorrencia(focos, territorio)
    q1, q2, q3 = risco["quartis"]
    assert q1 <= q2 <= q3
    classes = risco["classes"]
    assert [c["codigo"] for c in classes] == ["baixo", "moderado", "alto", "muito_alto"]
    grade = grade_qgis(focos, territorio)
    assert sum(c["celulas"] for c in classes) == len(grade["features"])     # toda célula com foco numa classe
    for c in classes:
        if c["celulas"]:
            assert c["geometria"].geom_type in ("Polygon", "MultiPolygon")   # dissolvido: um só polígono
    assert classes[0]["anos"][1] <= q1 < classes[3]["anos"][0]


def test_pasta_risco_de_recorrencia_no_pacote(municipio):
    import xml.dom.minidom
    nomes = [n for n in conteudo(municipio[2]).namelist() if n.endswith(".kml")]
    assert sorted(nomes) == [f"risco_de_recorrencia/risco_{i}_{c}_NOVO_PROGRESSO.kml" for i, c in
                             enumerate(["baixo", "moderado", "alto", "muito_alto"], start=1)]
    # P26: as 4 classes juntas num GeoJSON (só as classes com células)
    geo = json.loads(conteudo(municipio[2]).read("risco_de_recorrencia/risco_recorrencia_classes_NOVO_PROGRESSO.geojson"))
    assert 1 <= len(geo["features"]) <= 4
    assert {f["properties"]["classe"] for f in geo["features"]} <= {"baixo", "moderado", "alto", "muito_alto"}
    assert sum(f["properties"]["celulas"] for f in geo["features"]) == len(
        grade_qgis(municipio[0].focos, municipio[0].poligono)["features"])
    for nome in nomes:
        documento = xml.dom.minidom.parseString(conteudo(municipio[2]).read(nome))
        assert documento.getElementsByTagName("Document")
        assert len(documento.getElementsByTagName("Placemark")) <= 1         # uma feição dissolvida (ou nenhuma)


def test_cruzes_so_quando_ha_poucos_focos(municipio):
    diag = municipio[0]
    assert len(diag.focos) > MAX_CRUZES
    figura = mapa_densidade(diag.focos, diag.poligono, "X", 2023, 2024)
    assert any("não marcados" in t.get_text() for t in figura.texts)
    assert not any(c.get_offsets().shape[0] == len(diag.focos) for c in figura.axes[0].collections)
    fechar_figuras({"f": figura})
    poucos = diag.focos.head(50)
    figura = mapa_densidade(poucos, diag.poligono, "X", 2023, 2024)
    assert any("+ cinza = cada foco" in t.get_text() for t in figura.texts)
    fechar_figuras({"f": figura})


def test_fonte_de_imagem_aberta():
    """P22: imagem de fundo com licença aberta (Sentinel-2 cloudless, EOX), configurável em core/config.py."""
    from core import config
    assert "eox.at" in config.URL_SATELITE and "google" not in config.URL_SATELITE.lower()
    assert "CC BY-NC-SA 4.0" in config.CREDITO_SATELITE
    assert all(marcador in config.URL_SATELITE for marcador in ("{x}", "{y}", "{z}"))



def test_texto_dos_dados_nao_vira_formula_no_excel(tmp_path):
    """Texto que começa com "=" (ex.: nome vindo de uma base externa) fica como texto na planilha."""
    import openpyxl

    from core.texto import neutralizar_formulas

    caminho = tmp_path / "planilha.xlsx"
    with pd.ExcelWriter(caminho, engine="openpyxl") as planilha:
        pd.DataFrame({"nome": ['=HYPERLINK("http://exemplo.invalido","clique")', "Belém"]}).to_excel(planilha, index=False)
        neutralizar_formulas(planilha.book)
    celulas = [c for linha in openpyxl.load_workbook(caminho).active.iter_rows() for c in linha]
    assert [c.data_type for c in celulas] == ["s", "s", "s"]
    assert celulas[1].value.startswith("=HYPERLINK")
